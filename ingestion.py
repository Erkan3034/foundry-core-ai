"""
Foundry RAG Assistant - Ingestion Module
Belgeleri oku, parcala, embedding olustur, veritabanina kaydet.
"""

import io
import os
import logging
from pathlib import Path
from typing import List, Optional
from database import db
from embeddings import EmbeddingManager
from config import CONFIG

logger = logging.getLogger(__name__)


class DocumentIngestor:
    """Belge yukleme ve isleme pipeline'i."""

    SUPPORTED_EXTENSIONS = {'.txt', '.md', '.py', '.js', '.html', '.css',
                            '.json', '.xml', '.csv', '.log', '.rst',
                            '.pdf', '.docx', '.xlsx', '.pptx'}

    def __init__(self):
        self.chunk_size = CONFIG.chunk_size
        self.chunk_overlap = CONFIG.chunk_overlap
        self.embedding_manager = EmbeddingManager()

    def _read_file(self, file_path: Path) -> str:
        """Dosyayi oku (format otomatik algilanir, encoding otomatik tespit)."""
        ext = file_path.suffix.lower()
        reader = {
            '.pdf': self._read_pdf,
            '.docx': self._read_docx,
            '.xlsx': self._read_xlsx,
            '.pptx': self._read_pptx,
        }.get(ext)
        if reader:
            return reader(file_path)

        raw = file_path.read_bytes()

        # UTF-8 ile dene
        try:
            return raw.decode('utf-8')
        except UnicodeDecodeError:
            pass

        # BOM ile UTF-16 tespiti
        if raw[:2] in (b'\xff\xfe', b'\xfe\xff'):
            enc = 'utf-16-le' if raw[:2] == b'\xff\xfe' else 'utf-16-be'
            return raw.decode(enc).lstrip('\ufeff')

        # charset-normalizer ile encoding tespiti (ASCII, UTF-8, ISO-8859 vb.)
        try:
            from charset_normalizer import from_bytes
            result = from_bytes(raw).best()
            if result and result.encoding:
                return str(result)
        except Exception:
            pass

        # Turkce ve yaygin encoding'lerle dene
        for enc in ['windows-1254', 'iso-8859-9', 'latin-1', 'cp1252']:
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue

        # Hala basarisizsa latin-1 ile oku (kayip olabilir ama calisir)
        return raw.decode('latin-1', errors='replace')

    def _read_pdf(self, file_path: Path) -> str:
        import fitz
        text_parts = []
        with fitz.open(file_path) as doc:
            for page_num, page in enumerate(doc, 1):
                p_text = page.get_text().strip()
                if not p_text:
                    logger.warning(f"  [UYARI] PDF sayfa {page_num} metinsiz (görsel/OCR gerekebilir): {file_path.name}")
                    continue
                text_parts.append(f"--- SAYFA {page_num} ---\n{p_text}")
        return "\n\n".join(text_parts)

    def _read_docx(self, file_path: Path) -> str:
        from docx import Document
        doc = Document(file_path)
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    def _read_xlsx(self, file_path: Path) -> str:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        text_parts = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cleaned = [str(c).strip() if c is not None else "" for c in row]
                if any(cleaned):
                    rows.append(cleaned)
            if not rows:
                continue

            # Markdown Tablo Formatına Dönüştür
            table_lines = [f"### [Excel Sayfası: {sheet.title}]"]
            headers = rows[0]
            table_lines.append("| " + " | ".join(headers) + " |")
            table_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
            for r in rows[1:]:
                table_lines.append("| " + " | ".join(r) + " |")
            text_parts.append("\n".join(table_lines))
        wb.close()
        return "\n\n".join(text_parts)

    def _read_pptx(self, file_path: Path) -> str:
        from pptx import Presentation
        prs = Presentation(file_path)
        text_parts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                text_parts.append(f"--- SAYFA {slide_num} ---\n" + "\n".join(slide_texts))
        return "\n\n".join(text_parts)

    def _chunk_document(self, text: str) -> List[dict]:
        """Yapı-farkında (Structure-Aware) ve Parent-Child parçalama.

        Döndürür: [{'chunk_text': str, 'parent_chunk_text': str, 'section_title': str, 'page_number': int}]
        """
        import re
        if not text.strip():
            return []

        # Sayfa ve başlık bazlı bölümler
        lines = text.split("\n")
        sections = []
        current_section_lines = []
        current_title = "Genel"
        current_page = None

        for line in lines:
            # Sayfa kontrolü
            page_match = re.search(r'---\s*SAYFA\s*(\d+)\s*---', line, re.IGNORECASE)
            if page_match:
                current_page = int(page_match.group(1))

            # Başlık kontrolü (#, ##, Bölüm, Madde, vb.)
            header_match = re.match(r'^(#{1,6}\s+|Bölüm\s+\d+|Madde\s+\d+|###?\s+\[Excel Sayfası:)', line.strip(), re.IGNORECASE)
            if header_match and current_section_lines:
                sec_text = "\n".join(current_section_lines).strip()
                if sec_text:
                    sections.append({
                        "title": current_title,
                        "page_number": current_page,
                        "text": sec_text
                    })
                current_section_lines = []
                current_title = line.strip().lstrip('#').strip()

            current_section_lines.append(line)

        if current_section_lines:
            sec_text = "\n".join(current_section_lines).strip()
            if sec_text:
                sections.append({
                    "title": current_title,
                    "page_number": current_page,
                    "text": sec_text
                })

        # Parent-Child parçaları üret
        results = []
        for sec in sections:
            parent_text = sec["text"]
            # Eğer bölüm genişse (ör. > 400 karakter), küçük child chunk'lara böl
            if len(parent_text) > self.chunk_size:
                child_texts = self._chunk_text(parent_text)
                for child_t in child_texts:
                    results.append({
                        "chunk_text": child_t,
                        "parent_chunk_text": parent_text,
                        "section_title": sec["title"],
                        "page_number": sec["page_number"],
                    })
            else:
                results.append({
                    "chunk_text": parent_text,
                    "parent_chunk_text": parent_text,
                    "section_title": sec["title"],
                    "page_number": sec["page_number"],
                })

        return results

    def _chunk_text(self, text: str) -> List[str]:
        """Metni sabit/paragraf boyutlu parçalara böl."""
        if not text.strip():
            return []

        paragraphs = [p.strip() for p in text.split('\n\n') if p.strip()]

        chunks = []
        current_chunk = ""

        for paragraph in paragraphs:
            if len(paragraph) > self.chunk_size:
                if current_chunk:
                    chunks.append(current_chunk.strip())
                    current_chunk = ""

                sentences = paragraph.replace('. ', '.|').split('|')
                temp_chunk = ""
                for sentence in sentences:
                    if len(sentence) >= self.chunk_size:
                        # Cumle tek basina limiti asiyor (or. ayirac icermeyen
                        # uzun metin); sabit boyutlu dilimlere bol
                        if temp_chunk.strip():
                            chunks.append(temp_chunk.strip())
                            temp_chunk = ""
                        for j in range(0, len(sentence), self.chunk_size):
                            piece = sentence[j:j + self.chunk_size].strip()
                            if piece:
                                chunks.append(piece)
                    elif len(temp_chunk) + len(sentence) < self.chunk_size:
                        temp_chunk += sentence + " "
                    else:
                        if temp_chunk.strip():
                            chunks.append(temp_chunk.strip())
                        temp_chunk = sentence + " "
                if temp_chunk.strip():
                    chunks.append(temp_chunk.strip())
            else:
                if len(current_chunk) + len(paragraph) < self.chunk_size:
                    current_chunk += paragraph + "\n\n"
                else:
                    if current_chunk.strip():
                        chunks.append(current_chunk.strip())
                    current_chunk = paragraph + "\n\n"

        if current_chunk.strip():
            chunks.append(current_chunk.strip())

        if self.chunk_overlap > 0 and len(chunks) > 1:
            chunks = self._apply_overlap(chunks)

        return chunks

    def _estimate_token_count(self, text: str) -> int:
        """Metnin tahmini token sayisini hesapla (Turkce sub-word katsayisi ile)."""
        words = text.split()
        if not words:
            return 0
        return max(1, int(len(words) * 1.5))

    def _apply_overlap(self, chunks: List[str]) -> List[str]:
        """Parcalar arasi sinir duyarli (cumle/kelime kesilmeyen) overlap uygula."""
        if not chunks or self.chunk_overlap <= 0:
            return chunks

        overlapped = [chunks[0]]
        for i in range(1, len(chunks)):
            prev_chunk = chunks[i - 1]
            budget = min(self.chunk_overlap, self.chunk_size - len(chunks[i]) - 1)
            if budget <= 0:
                overlapped.append(chunks[i])
                continue

            raw_slice = prev_chunk[-budget:]
            space_idx = raw_slice.find(' ')
            if space_idx != -1 and space_idx < len(raw_slice) - 1:
                clean_overlap = raw_slice[space_idx + 1:]
            else:
                clean_overlap = raw_slice

            if clean_overlap.strip():
                overlapped.append(clean_overlap.strip() + "\n" + chunks[i])
            else:
                overlapped.append(chunks[i])

        return overlapped

    def ingest_file(self, file_path: Path, source_name: Optional[str] = None, force: bool = False) -> dict:
        """Tek bir dosyayi isle."""
        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(f"Dosya bulunamadi: {file_path}")

        if file_path.suffix.lower() not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Desteklenmeyen dosya formati: {file_path.suffix}")

        if not force and db.is_file_ingested(str(file_path)):
            logger.info(f"Atlandi (zaten islenmis): {file_path}")
            return {"document_id": None, "chunks": 0, "status": "skipped"}

        source = source_name or file_path.name
        content = self._read_file(file_path)

        if not content.strip():
            logger.warning(f"Bos dosya atlandi: {file_path}")
            return {"document_id": None, "chunks": 0, "status": "empty"}

        doc_id = db.add_document(
            source=source,
            file_path=str(file_path),
            content=content
        )

        chunk_objs = self._chunk_document(content)
        logger.info(f"{source}: {len(chunk_objs)} parca olusturuldu (Structure-Aware / Parent-Child)")

        if not chunk_objs:
            return {"document_id": doc_id, "chunks": 0, "status": "no_chunks"}

        chunk_texts = [c["chunk_text"] for c in chunk_objs]
        logger.info(f"Embedding olusturuluyor ({len(chunk_objs)} parca)...")
        
        # Batching: 16'lik gruplar halinde bellek tasmalarini onle
        embeddings = []
        BATCH_SIZE = 16
        for b in range(0, len(chunk_texts), BATCH_SIZE):
            batch_texts = chunk_texts[b:b + BATCH_SIZE]
            batch_embeddings = self.embedding_manager.embed_batch(batch_texts)
            embeddings.extend(batch_embeddings)

        db.add_chunks_bulk(doc_id, [
            {
                "chunk_index": i,
                "chunk_text": obj["chunk_text"],
                "parent_chunk_text": obj.get("parent_chunk_text"),
                "section_title": obj.get("section_title"),
                "page_number": obj.get("page_number"),
                "embedding": embedding,
                "token_count": self._estimate_token_count(obj["chunk_text"]),
            }
            for i, (obj, embedding) in enumerate(zip(chunk_objs, embeddings))
        ])

        logger.info(f"{source} islendi: {len(chunks)} parca, {len(embeddings)} embedding")

        try:
            from retriever import Retriever
            Retriever.invalidate_cache()
        except Exception:
            pass

        return {
            "document_id": doc_id,
            "chunks": len(chunks),
            "status": "success"
        }

    def ingest_directory(self, directory: Path, recursive: bool = True, force: bool = False) -> List[dict]:
        """Bir dizindeki tum desteklenen dosyalari isle."""
        directory = Path(directory)

        if not directory.exists():
            raise FileNotFoundError(f"Dizin bulunamadi: {directory}")

        results = []
        pattern = "**/*" if recursive else "*"

        for file_path in directory.glob(pattern):
            if file_path.is_file() and file_path.suffix.lower() in self.SUPPORTED_EXTENSIONS:
                try:
                    result = self.ingest_file(file_path, force=force)
                    results.append(result)
                except Exception as e:
                    logger.error(f"Hata ({file_path}): {e}")
                    results.append({
                        "file": str(file_path),
                        "status": "error",
                        "error": str(e)
                    })

        return results

    def ingest_text(self, text: str, source_name: str = "manual") -> dict:
        """Dogrudan metin icerigi isle."""
        doc_id = db.add_document(
            source=source_name,
            file_path="inline",
            content=text
        )

        chunks = self._chunk_text(text)
        embeddings = self.embedding_manager.embed_batch(chunks)

        for i, (chunk_text, embedding) in enumerate(zip(chunks, embeddings)):
            db.add_chunk(
                document_id=doc_id,
                chunk_index=i,
                chunk_text=chunk_text,
                embedding=embedding
            )

        return {
            "document_id": doc_id,
            "chunks": len(chunks),
            "status": "success"
        }

    def shutdown(self):
        """Kaynaklari temizle."""
        self.embedding_manager.shutdown()


def run_ingestion(documents_dir: Optional[Path] = None, force: bool = False):
    """Ana ingestion fonksiyonu. CLI'dan cagrılır."""
    docs_dir = documents_dir or CONFIG.docs_path

    logger.info(f"Ingestion baslatiliyor: {docs_dir}")

    ingestor = DocumentIngestor()

    try:
        ingestor.embedding_manager.initialize()
        results = ingestor.ingest_directory(docs_dir, force=force)

        success_count = sum(1 for r in results if r.get("status") == "success")
        logger.info(f"Ingestion tamamlandi: {success_count}/{len(results)} basarili")

        stats = db.get_stats()
        logger.info(f"Veritabani: {stats['documents']} belge, {stats['chunks']} parca, {stats['embedded_chunks']} embedding")

    finally:
        ingestor.shutdown()


if __name__ == "__main__":
    import sys
    logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')

    docs_path = sys.argv[1] if len(sys.argv) > 1 else CONFIG.docs_path
    run_ingestion(Path(docs_path))
